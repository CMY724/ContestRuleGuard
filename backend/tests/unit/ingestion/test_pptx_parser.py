# backend/tests/unit/ingestion/test_pptx_parser.py
from io import BytesIO
from uuid import uuid4

from pptx import Presentation
from pptx.util import Inches

from contest_rule_guard.ingestion.models import BlockKind, ParseContext, UnitKind
from contest_rule_guard.ingestion.parsers.pptx import PptxParser


def make_pptx() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "匿名材料不得出现学校名称"
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(6), Inches(1))
    table_shape.table.cell(0, 0).text = "PPT"
    table_shape.table.cell(0, 1).text = "不超过20页"
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_pptx_parser_preserves_slide_and_shape_locations() -> None:
    content = make_pptx()
    context = ParseContext.from_upload(
        document_id=uuid4(),
        filename="答辩规则.pptx",
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        content=content,
    )
    parsed = PptxParser().parse(content, context)
    assert parsed.units[0].kind is UnitKind.SLIDE
    assert parsed.units[0].index == 1
    texts = [block.text for block in parsed.units[0].blocks]
    assert "匿名材料不得出现学校名称" in texts
    assert "PPT" in texts
    assert "不超过20页" in texts
    assert parsed.units[0].blocks[0].source_path == "slide[1].shape[0].paragraph[0]"
    assert parsed.units[0].blocks[1].kind is BlockKind.TABLE_CELL
    assert parsed.units[0].blocks[1].source_path.startswith("slide[1].shape[1].table")
    assert parsed.units[0].blocks[0].bbox is not None
