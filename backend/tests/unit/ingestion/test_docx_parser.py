# backend/tests/unit/ingestion/test_docx_parser.py
from io import BytesIO
from uuid import uuid4

from docx import Document

from contest_rule_guard.ingestion.models import BlockKind, ParseContext, UnitKind
from contest_rule_guard.ingestion.parsers.docx import DocxParser


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("提交规则")
    document.add_paragraph("团队人数：1至3人")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "赛道"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "AI+OPC"
    table.cell(1, 1).text = "演示脚本"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_parser_captures_paragraph_and_table_location() -> None:
    content = make_docx()
    context = ParseContext.from_upload(
        document_id=uuid4(),
        filename="rules.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=content,
    )
    parsed = DocxParser().parse(content, context)
    assert parsed.units[0].kind is UnitKind.FLOW
    kinds = [block.kind for block in parsed.units[0].blocks]
    texts = [block.text for block in parsed.units[0].blocks]
    assert BlockKind.PARAGRAPH in kinds
    assert BlockKind.TABLE_CELL in kinds
    assert "提交规则" in texts
    assert "团队人数：1至3人" in texts
    assert "赛道" in texts
    assert "AI+OPC" in texts
    assert "演示脚本" in texts
