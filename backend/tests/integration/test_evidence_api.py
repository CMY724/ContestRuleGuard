from uuid import uuid4

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"  # noqa: E501


def test_create_and_list_evidence(client) -> None:
    # Create a project
    response = client.post(
        "/api/projects",
        json={
            "name": "evidence-int",
            "competition_name": "?????????",
            "competition_year": 2026,
            "track": "AI",
            "target_stage": "national",
        },
    )
    assert response.status_code == 201, response.text
    project_id = response.json()["id"]

    # Upload a document
    from io import BytesIO

    from docx import Document
    docx = Document()
    docx.add_paragraph("????????????????????")
    buffer = BytesIO()
    docx.save(buffer)
    buffer.seek(0)

    upload = client.post(
        f"/api/projects/{project_id}/documents",
        data={"source_tier": "primary", "stage": "school"},
        files={"file": ("rule.docx", buffer, DOCX_MIME)},
    )
    assert upload.status_code == 201, upload.text
    doc = upload.json()
    doc_id = doc["id"]
    block = doc["units"][0]["blocks"][0]

    # Create evidence
    span_data = {
        "document_id": doc_id,
        "block_id": block["id"],
        "field_path": "/payload/eligibility",
        "quote": block["text"],
        "start_char": 0,
        "end_char": len(block["text"]),
        "source_tier": "primary",
        "stage": "school",
        "document_sha256": doc["content_sha256"],
    }
    evidence_response = client.post(
        f"/api/projects/{project_id}/documents/{doc_id}/evidence",
        json=span_data,
    )
    assert evidence_response.status_code == 201, evidence_response.text
    evidence = evidence_response.json()
    assert evidence["field_path"] == "/payload/eligibility"
    assert evidence["quote"] == block["text"]

    # List evidence
    listed = client.get(
        f"/api/projects/{project_id}/documents/{doc_id}/evidence",
    )
    assert listed.status_code == 200
    assert len(listed.json()) == 1
    assert listed.json()[0]["id"] == evidence["id"]

    # Get specific evidence
    get_resp = client.get(
        f"/api/projects/{project_id}/documents/{doc_id}/evidence/{evidence['id']}",
    )
    assert get_resp.status_code == 200
    assert get_resp.json()["id"] == evidence["id"]

    # GET with non-existent evidence returns 404
    fake_id = str(uuid4())
    not_found = client.get(
        f"/api/projects/{project_id}/documents/{doc_id}/evidence/{fake_id}",
    )
    assert not_found.status_code == 404


def test_evidence_with_invalid_block_returns_422(client) -> None:
    response = client.post(
        "/api/projects",
        json={
            "name": "evidence-422",
            "competition_name": "?????????",
            "competition_year": 2026,
            "track": "AI",
            "target_stage": "national",
        },
    )
    project_id = response.json()["id"]

    from io import BytesIO

    from docx import Document
    docx = Document()
    docx.add_paragraph("????????????????????")
    buffer = BytesIO()
    docx.save(buffer)
    buffer.seek(0)

    upload = client.post(
        f"/api/projects/{project_id}/documents",
        data={"source_tier": "primary", "stage": "school"},
        files={"file": ("rule.docx", buffer, DOCX_MIME)},
    )
    doc = upload.json()

    # Try creating evidence with non-existent block_id
    span_data = {
        "document_id": doc["id"],
        "block_id": str(uuid4()),
        "field_path": "/payload/test",
        "quote": "test text",
        "start_char": 0,
        "end_char": 9,
        "source_tier": "primary",
        "stage": "school",
        "document_sha256": doc["content_sha256"],
    }
    bad = client.post(
        f"/api/projects/{project_id}/documents/{doc['id']}/evidence",
        json=span_data,
    )
    assert bad.status_code == 422, bad.text


def test_evidence_for_nonexistent_document_returns_404(client) -> None:
    response = client.post(
        "/api/projects",
        json={
            "name": "evidence-404",
            "competition_name": "?????????",
            "competition_year": 2026,
            "track": "AI",
            "target_stage": "national",
        },
    )
    project_id = response.json()["id"]
    fake_doc_id = str(uuid4())

    # List evidence for nonexistent doc
    listed = client.get(
        f"/api/projects/{project_id}/documents/{fake_doc_id}/evidence",
    )
    assert listed.status_code == 404

    # Create evidence for nonexistent doc
    span_data = {
        "document_id": fake_doc_id,
        "block_id": str(uuid4()),
        "field_path": "/payload/test",
        "quote": "test",
        "start_char": 0,
        "end_char": 4,
        "source_tier": "primary",
        "stage": "school",
        "document_sha256": "a" * 64,
    }
    create_resp = client.post(
        f"/api/projects/{project_id}/documents/{fake_doc_id}/evidence",
        json=span_data,
    )
    assert create_resp.status_code == 404
