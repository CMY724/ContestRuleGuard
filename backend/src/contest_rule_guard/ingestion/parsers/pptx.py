from io import BytesIO

from pptx import Presentation
from pptx.shapes.base import BaseShape

from contest_rule_guard.ingestion.models import (
    BlockKind,
    BoundingBox,
    DocumentUnit,
    NormalizedDocument,
    ParseContext,
    TextBlock,
    UnitKind,
)
from contest_rule_guard.ingestion.ports import DocumentParser


def _zero_dimension(name: str) -> int:
    raise ValueError(f"invalid pptx: zero EMU value for {name}")


def _shape_bbox(shape: BaseShape, width: int, height: int) -> BoundingBox:
    left = int(shape.left)
    top = int(shape.top)
    right = left + int(shape.width)
    bottom = top + int(shape.height)
    return BoundingBox(
        left=left / width if width else _zero_dimension("slide width"),
        top=top / height if height else _zero_dimension("slide height"),
        right=right / width if width else _zero_dimension("slide width"),
        bottom=bottom / height if height else _zero_dimension("slide height"),
    )


class PptxParser(DocumentParser):
    media_types = frozenset(
        {"application/vnd.openxmlformats-officedocument.presentationml.presentation"}
    )
    extensions = frozenset({".pptx"})

    def parse(self, content: bytes, context: ParseContext) -> NormalizedDocument:
        deck = Presentation(BytesIO(content))
        sw = deck.slide_width or 0
        sh = deck.slide_height or 0
        units: list[DocumentUnit] = []
        for slide_number, slide in enumerate(deck.slides, start=1):
            blocks: list[TextBlock] = []
            for shape_index, shape in enumerate(slide.shapes):
                bbox = _shape_bbox(shape, sw, sh)
                if shape.has_table:
                    table = shape.table
                    for row_index, row in enumerate(table.rows):
                        for cell_index, cell in enumerate(row.cells):
                            text = " ".join(cell.text.split())
                            if text:
                                prefix = f"slide[{slide_number}].shape[{shape_index}]"
                                suffix = (
                                    f".table.row[{row_index}]"
                                    f".cell[{cell_index}]"
                                )
                                path = f"{prefix}{suffix}"
                                blocks.append(
                                    TextBlock.build(
                                        document_id=context.document_id,
                                        unit_index=slide_number,
                                        block_index=len(blocks),
                                        source_path=path,
                                        text=text,
                                        kind=BlockKind.TABLE_CELL,
                                        bbox=bbox,
                                    )
                                )
                elif shape.has_text_frame:
                    for par_i, paragraph in enumerate(
                        shape.text_frame.paragraphs
                    ):
                        text = " ".join(paragraph.text.split())
                        if text:
                            prefix = f"slide[{slide_number}].shape[{shape_index}]"
                            suffix = f".paragraph[{par_i}]"
                            path = f"{prefix}{suffix}"
                            blocks.append(
                                TextBlock.build(
                                    document_id=context.document_id,
                                    unit_index=slide_number,
                                    block_index=len(blocks),
                                    source_path=path,
                                    text=text,
                                    kind=BlockKind.TEXT_BOX,
                                    bbox=bbox,
                                )
                            )
            units.append(
                DocumentUnit(
                    index=slide_number,
                    kind=UnitKind.SLIDE,
                    width=sw,
                    height=sh,
                    blocks=blocks,
                )
            )
        return NormalizedDocument(
            id=context.document_id,
            filename=context.filename,
            media_type=context.media_type,
            content_sha256=context.content_sha256,
            units=units,
        )
