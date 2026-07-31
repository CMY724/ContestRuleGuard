# backend/src/contest_rule_guard/ingestion/parsers/docx.py
from io import BytesIO

from docx import Document

from contest_rule_guard.ingestion.models import (
    BlockKind,
    DocumentUnit,
    NormalizedDocument,
    ParseContext,
    TextBlock,
    UnitKind,
)
from contest_rule_guard.ingestion.ports import DocumentParser


class DocxParser(DocumentParser):
    media_types = frozenset(
        {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    )
    extensions = frozenset({".docx"})

    def parse(self, content: bytes, context: ParseContext) -> NormalizedDocument:
        source = Document(BytesIO(content))
        blocks: list[TextBlock] = []
        paragraph_index = 0
        table_index = 0

        for element in source.element.body.iterchildren():
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "p":
                text = source.paragraphs[paragraph_index].text.strip()
                if text:
                    blocks.append(
                        TextBlock.build(
                            document_id=context.document_id,
                            unit_index=1,
                            block_index=len(blocks),
                            source_path=f"paragraph[{paragraph_index}]",
                            text=text,
                            kind=BlockKind.PARAGRAPH,
                        )
                    )
                paragraph_index += 1
            elif tag == "tbl":
                table = source.tables[table_index]
                for row_index, row in enumerate(table.rows):
                    for cell_index, cell in enumerate(row.cells):
                        text = " ".join(cell.text.split())
                        if text:
                            blocks.append(
                                TextBlock.build(
                                    document_id=context.document_id,
                                    unit_index=1,
                                    block_index=len(blocks),
                                    source_path=f"table[{table_index}].row[{row_index}].cell[{cell_index}]",
                                    text=text,
                                    kind=BlockKind.TABLE_CELL,
                                )
                            )
                table_index += 1

        return NormalizedDocument(
            id=context.document_id,
            filename=context.filename,
            media_type=context.media_type,
            content_sha256=context.content_sha256,
            units=[DocumentUnit(index=1, kind=UnitKind.FLOW, blocks=blocks)],
        )
